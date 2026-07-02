import ast
import re
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

workspace = Path.cwd()
source = workspace / 'build_feedback_workbook.mjs'
output_dir = workspace / 'outputs' / 'feedback_dataset'
output_dir.mkdir(parents=True, exist_ok=True)
output = output_dir / 'Customer_Feedback_Data.xlsx'

headers = ['Date', 'Customer', 'Feedback Type', 'Rating (1-5)', 'Status', 'Follow-Up', 'Comments']
rows = []
for line in source.read_text(encoding='utf-8').splitlines():
    stripped = line.strip()
    if stripped.startswith('["2025-'):
        rows.append(ast.literal_eval(stripped.rstrip(',')))

wb = Workbook()
ws = wb.active
ws.title = 'Feedback Data'
summary = wb.create_sheet('Summary')

ws.append(headers)
for row in rows:
    ws.append(row)

for cell in ws[1]:
    cell.fill = PatternFill('solid', fgColor='1F4E78')
    cell.font = Font(color='FFFFFF', bold=True)
    cell.alignment = Alignment(horizontal='center')

thin_blue = Side(style='thin', color='B7C9D9')
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
    for cell in row:
        cell.border = Border(bottom=thin_blue)
        cell.alignment = Alignment(vertical='top')

for cell in ws['A'][1:]:
    cell.number_format = 'yyyy-mm-dd'
for cell in ws['D'][1:]:
    cell.number_format = '0'
    cell.alignment = Alignment(horizontal='center')
for col in ['E', 'F']:
    for cell in ws[col][1:]:
        cell.alignment = Alignment(horizontal='center')

widths = {'A': 13, 'B': 14, 'C': 20, 'D': 13, 'E': 15, 'F': 12, 'G': 46}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

ws.freeze_panes = 'A2'
tab = Table(displayName='FeedbackTable', ref=f'A1:G{ws.max_row}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
ws.add_table(tab)

rating_validation = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=False)
status_validation = DataValidation(type='list', formula1='"Resolved,Not Resolved"')
follow_validation = DataValidation(type='list', formula1='"Yes,No"')
ws.add_data_validation(rating_validation)
ws.add_data_validation(status_validation)
ws.add_data_validation(follow_validation)
rating_validation.add(f'D2:D{ws.max_row}')
status_validation.add(f'E2:E{ws.max_row}')
follow_validation.add(f'F2:F{ws.max_row}')

summary.merge_cells('A1:D1')
summary['A1'] = 'Customer Feedback Summary'
summary['A1'].fill = PatternFill('solid', fgColor='1F4E78')
summary['A1'].font = Font(color='FFFFFF', bold=True, size=16)
summary['A1'].alignment = Alignment(horizontal='center')

summary_rows = [
    ['Metric', 'Value'],
    ['Total Feedback', f'=COUNTA(\'Feedback Data\'!B2:B{ws.max_row})'],
    ['Average Rating', f'=AVERAGE(\'Feedback Data\'!D2:D{ws.max_row})'],
    ['Resolved', f'=COUNTIF(\'Feedback Data\'!E2:E{ws.max_row},"Resolved")'],
    ['Not Resolved', f'=COUNTIF(\'Feedback Data\'!E2:E{ws.max_row},"Not Resolved")'],
    ['Follow-Up Needed', f'=COUNTIF(\'Feedback Data\'!F2:F{ws.max_row},"Yes")'],
]
for r, row in enumerate(summary_rows, start=3):
    for c, value in enumerate(row, start=1):
        summary.cell(r, c, value)

feedback_types = ['Service', 'Loan Process', 'Online Banking', 'ATM', 'App', 'Mobile App', 'Customer Service', 'Branch Experience', 'Debit Card', 'POS Transaction', 'Account Opening', 'Internet Banking']
summary['D3'] = 'Feedback Type'
summary['E3'] = 'Count'
for idx, item in enumerate(feedback_types, start=4):
    summary.cell(idx, 4, item)
    summary.cell(idx, 5, f'=COUNTIF(\'Feedback Data\'!C2:C{ws.max_row},D{idx})')

for rng in ['A3:B8', f'D3:E{3 + len(feedback_types)}']:
    for row in summary[rng]:
        for cell in row:
            cell.border = Border(left=thin_blue, right=thin_blue, top=thin_blue, bottom=thin_blue)
            cell.alignment = Alignment(vertical='center')
for cell in list(summary['A3:B3'][0]) + list(summary['D3:E3'][0]):
    cell.fill = PatternFill('solid', fgColor='D9EAF7')
    cell.font = Font(bold=True, color='12344D')
summary['B5'].number_format = '0.0'
for col, width in {'A': 22, 'B': 16, 'D': 22, 'E': 12}.items():
    summary.column_dimensions[col].width = width

wb.save(output)

check = load_workbook(output, data_only=False)
assert check['Feedback Data'].max_row == len(rows) + 1
assert check['Feedback Data']['A1'].value == 'Date'
assert check['Summary']['B4'].value.startswith('=COUNTA')
print(output)
print(f'Rows: {len(rows)}')
